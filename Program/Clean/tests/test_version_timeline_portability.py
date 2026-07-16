import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


CLEAN_DIR = Path(__file__).resolve().parents[1]
SCRIPT = CLEAN_DIR / "build_version_timeline.py"


class VersionTimelinePortabilityTests(unittest.TestCase):
    def test_help_exposes_explicit_inputs_and_output(self):
        result = subprocess.run(
            [sys.executable, str(SCRIPT), "--help"],
            check=False,
            capture_output=True,
            text=True,
        )

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("--catalog-csv", result.stdout)
        self.assertIn("--official-xlsx", result.stdout)
        self.assertIn("--anchor-dir", result.stdout)
        self.assertIn("--coverage-manifest", result.stdout)
        self.assertIn("--cover-dates-csv", result.stdout)
        self.assertIn("--date-corrections-csv", result.stdout)
        self.assertIn("--frozen-input-dir", result.stdout)
        self.assertIn("--output-dir", result.stdout)

    def test_source_has_no_legacy_result_or_release_path(self):
        source = SCRIPT.read_text(encoding="utf-8")

        for forbidden in ("03_Results", "04_Release", "01_RawData"):
            self.assertNotIn(forbidden, source)

    def test_cli_builds_timeline_from_injected_temporary_paths(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "inputs" / "catalog.csv"
            workbook = root / "inputs" / "official.xlsx"
            anchor = root / "anchor"
            output = root / "output"
            catalog.parent.mkdir()
            category = anchor / "污染物限量"
            category.mkdir(parents=True)
            (category / "GB 2762-2022.pdf").touch()

            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerows(
                    [
                        {
                            "code": "GB",
                            "number": "2762",
                            "year": "2017",
                            "title": "食品安全国家标准 食品中污染物限量",
                            "path": "legacy/GB 2762-2017.pdf",
                        },
                        {
                            "code": "GB",
                            "number": "2762",
                            "year": "2022",
                            "title": "食品安全国家标准 食品中污染物限量",
                            "path": "anchor/污染物限量/GB 2762-2022.pdf",
                        },
                    ]
                )

            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            sheet.append(
                [
                    "GB 2762-2017",
                    "食品安全国家标准 食品中污染物限量",
                    "2017-03-17",
                    "2017-09-17",
                    "污染物限",
                ]
            )
            sheet.append(
                [
                    "GB 2762-2022",
                    "食品安全国家标准 食品中污染物限量",
                    "2022-06-30",
                    "2023-06-30",
                    "污染物限",
                ]
            )
            book.save(workbook)
            book.close()

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--output-dir",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            with (output / "standard_versions_long.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                long_rows = list(csv.DictReader(handle))
            with (output / "standard_timeline.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                timeline_rows = list(csv.DictReader(handle))

            self.assertEqual([row["year"] for row in long_rows], ["2017", "2022"])
            self.assertEqual(long_rows[1]["发布日期"], "2022-06-30")
            self.assertEqual(long_rows[1]["在新目录"], "1")
            self.assertEqual(long_rows[0]["path"], "GB 2762-2017.pdf")
            self.assertNotIn("/", long_rows[0]["path"])
            self.assertEqual(timeline_rows[0]["版本年份"], "2017;2022")
            self.assertEqual(timeline_rows[0]["卫健委类目"], "污染物限量")

            second_output = root / "second-output"
            second_result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--output-dir",
                    str(second_output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(second_result.returncode, 0, second_result.stderr)
            for filename in ("standard_versions_long.csv", "standard_timeline.csv"):
                self.assertNotIn(b"\r\n", (output / filename).read_bytes())
                self.assertEqual(
                    (output / filename).read_bytes(),
                    (second_output / filename).read_bytes(),
                )

    def test_reverse_chronology_fails_closed(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "catalog.csv"
            workbook = root / "official.xlsx"
            anchor = root / "anchor"
            output = root / "output"
            category = anchor / "食品添加剂"
            category.mkdir(parents=True)
            (category / "GB 1886.5-2015 食品添加剂 硝酸钠.pdf").touch()

            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "code": "GB",
                        "number": "1886.5",
                        "year": "2015",
                        "title": "食品安全国家标准 食品添加剂 硝酸钠",
                        "path": (
                            "anchor/食品添加剂/"
                            "GB 1886.5-2015 食品添加剂 硝酸钠.pdf"
                        ),
                    }
                )

            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            sheet.append(
                [
                    "GB 1886.5-2015",
                    "食品安全国家标准 食品添加剂 硝酸钠",
                    "2015-09-22",
                    "2015-03-22",
                    "食品添加剂",
                ]
            )
            book.save(workbook)
            book.close()

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--output-dir",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("issue date is later than effective date", result.stderr)

    def test_gb_1886_5_source_correction_restores_chronology(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "catalog.csv"
            workbook = root / "official.xlsx"
            corrections = root / "official_date_corrections.csv"
            anchor = root / "anchor"
            output = root / "output"
            category = anchor / "食品添加剂"
            category.mkdir(parents=True)
            (category / "GB 1886.5-2015 食品添加剂 硝酸钠.pdf").touch()

            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "code": "GB",
                        "number": "1886.5",
                        "year": "2015",
                        "title": "食品安全国家标准 食品添加剂 硝酸钠",
                        "path": (
                            "anchor/食品添加剂/"
                            "GB 1886.5-2015 食品添加剂 硝酸钠.pdf"
                        ),
                    }
                )

            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            sheet.append(
                [
                    "GB 1886.5-2015",
                    "食品安全国家标准 食品添加剂 硝酸钠",
                    "2015-09-22",
                    "2015-03-22",
                    "食品添加剂",
                ]
            )
            book.save(workbook)
            book.close()

            correction_fields = [
                "standard",
                "version_year",
                "event_type",
                "observed_issue_date",
                "observed_effective_date",
                "corrected_issue_date",
                "corrected_effective_date",
                "evidence_url",
                "evidence_file",
                "evidence_sha256",
                "verification_note",
            ]
            with corrections.open(
                "w", encoding="utf-8-sig", newline=""
            ) as handle:
                writer = csv.DictWriter(handle, fieldnames=correction_fields)
                writer.writeheader()
                writer.writerow(
                    {
                        "standard": "GB 1886.5",
                        "version_year": "2015",
                        "event_type": "official_version",
                        "observed_issue_date": "2015-09-22",
                        "observed_effective_date": "2015-03-22",
                        "corrected_issue_date": "2015-09-22",
                        "corrected_effective_date": "2016-03-22",
                        "evidence_url": (
                            "https://www.nhc.gov.cn/ewebeditor/uploadfile/"
                            "2015/10/20151030162720807.zip"
                        ),
                        "evidence_file": (
                            "GB 1886.5-2015 食品添加剂 硝酸钠.pdf"
                        ),
                        "evidence_sha256": (
                            "7f8806254dbe801d36f0b0c842db579c3a71e4caa3d7b255f"
                            "6d587ada5fd668a"
                        ),
                        "verification_note": (
                            "cover prints 2015-09-22 issue and 2016-03-22 effective"
                        ),
                    }
                )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--date-corrections-csv",
                    str(corrections),
                    "--output-dir",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            with (output / "standard_versions_long.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["发布日期"], "2015-09-22")
            self.assertEqual(rows[0]["实施日期"], "2016-03-22")

    def test_frozen_replay_updates_long_and_wide_then_rejects_stale_old_values(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            frozen = root / "frozen"
            output = root / "output"
            corrections = root / "official_date_corrections.csv"
            frozen.mkdir()

            long_columns = [
                "std_id",
                "number",
                "event_type",
                "year",
                "amend_no",
                "发布日期",
                "实施日期",
                "卫健委类目",
                "官方分类",
                "title",
                "在新目录",
                "path",
            ]
            with (frozen / "standard_versions_long.csv").open(
                "w", encoding="utf-8-sig", newline=""
            ) as handle:
                writer = csv.DictWriter(handle, fieldnames=long_columns)
                writer.writeheader()
                writer.writerow(
                    {
                        "std_id": "GB 1886.5",
                        "number": "1886.5",
                        "event_type": "正式版",
                        "year": "2015",
                        "amend_no": "",
                        "发布日期": "2015-09-22",
                        "实施日期": "2015-03-22",
                        "卫健委类目": "食品添加剂",
                        "官方分类": "食品添加剂",
                        "title": "食品添加剂 硝酸钠",
                        "在新目录": "1",
                        "path": "GB 1886.5-2015 食品添加剂 硝酸钠.pdf",
                    }
                )

            wide_columns = [
                "std_id",
                "title",
                "卫健委类目",
                "n_正式版本",
                "版本年份",
                "首版年",
                "现行版年",
                "现行版发布日期",
                "现行版实施日期",
                "n_修改单",
                "修改单年份",
                "有历史沿革",
            ]
            with (frozen / "standard_timeline.csv").open(
                "w", encoding="utf-8-sig", newline=""
            ) as handle:
                writer = csv.DictWriter(handle, fieldnames=wide_columns)
                writer.writeheader()
                writer.writerow(
                    {
                        "std_id": "GB 1886.5",
                        "title": "食品添加剂 硝酸钠",
                        "卫健委类目": "食品添加剂",
                        "n_正式版本": "1",
                        "版本年份": "2015",
                        "首版年": "2015",
                        "现行版年": "2015",
                        "现行版发布日期": "2015-09-22",
                        "现行版实施日期": "2015-03-22",
                        "n_修改单": "0",
                        "修改单年份": "",
                        "有历史沿革": "0",
                    }
                )

            correction_columns = [
                "standard",
                "version_year",
                "event_type",
                "observed_issue_date",
                "observed_effective_date",
                "corrected_issue_date",
                "corrected_effective_date",
                "evidence_url",
                "evidence_file",
                "evidence_sha256",
                "verification_note",
            ]
            with corrections.open(
                "w", encoding="utf-8-sig", newline=""
            ) as handle:
                writer = csv.DictWriter(handle, fieldnames=correction_columns)
                writer.writeheader()
                writer.writerow(
                    {
                        "standard": "GB 1886.5",
                        "version_year": "2015",
                        "event_type": "official_version",
                        "observed_issue_date": "2015-09-22",
                        "observed_effective_date": "2015-03-22",
                        "corrected_issue_date": "2015-09-22",
                        "corrected_effective_date": "2016-03-22",
                        "evidence_url": "https://www.nhc.gov.cn/official.zip",
                        "evidence_file": "GB 1886.5-2015.pdf",
                        "evidence_sha256": "a" * 64,
                        "verification_note": "verified against the cover",
                    }
                )

            command = [
                sys.executable,
                str(SCRIPT),
                "--frozen-input-dir",
                str(frozen),
                "--date-corrections-csv",
                str(corrections),
                "--output-dir",
                str(output),
            ]
            result = subprocess.run(
                command, check=False, capture_output=True, text=True
            )
            self.assertEqual(result.returncode, 0, result.stderr)

            with (output / "standard_versions_long.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                long_rows = list(csv.DictReader(handle))
            with (output / "standard_timeline.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                wide_rows = list(csv.DictReader(handle))
            self.assertEqual(long_rows[0]["实施日期"], "2016-03-22")
            self.assertEqual(wide_rows[0]["现行版实施日期"], "2016-03-22")
            before = {
                path.name: path.read_bytes()
                for path in (
                    output / "standard_versions_long.csv",
                    output / "standard_timeline.csv",
                )
            }

            stale_command = [
                str(output) if value == str(frozen) else value for value in command
            ]
            stale = subprocess.run(
                stale_command, check=False, capture_output=True, text=True
            )
            self.assertNotEqual(stale.returncode, 0)
            self.assertIn("must match exactly one declared old value", stale.stderr)
            after = {
                path.name: path.read_bytes()
                for path in (
                    output / "standard_versions_long.csv",
                    output / "standard_timeline.csv",
                )
            }
            self.assertEqual(after, before)

    def test_coverage_manifest_supplements_missing_veterinary_sources(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "catalog.csv"
            workbook = root / "official.xlsx"
            anchor = root / "anchor"
            manifest = root / "source_manifest.csv"
            output = root / "output"
            category = anchor / "污染物限量"
            category.mkdir(parents=True)
            (category / "GB 2762-2022.pdf").touch()

            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "code": "GB",
                        "number": "2762",
                        "year": "2022",
                        "title": "食品安全国家标准 食品中污染物限量",
                        "path": "anchor/污染物限量/GB 2762-2022.pdf",
                    }
                )

            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            sheet.append(
                [
                    "GB 2762-2022",
                    "食品安全国家标准 食品中污染物限量",
                    "2022-06-30",
                    "2023-06-30",
                    "污染物限量",
                ]
            )
            sheet.append(
                [
                    "GB 31650-2019",
                    "食品安全国家标准 食品中兽药最大残留限量",
                    "2019-09-06",
                    "2020-04-01",
                    "兽药残留限量",
                ]
            )
            sheet.append(
                [
                    "GB 31650.1-2022",
                    "食品安全国家标准 食品中兽药最大残留限量增补",
                    "2022-09-20",
                    "2023-02-01",
                    "兽药残留限量",
                ]
            )
            book.save(workbook)
            book.close()

            fields = [
                "source_id",
                "family",
                "standard",
                "version_year",
                "source_role",
                "release_status",
                "source_filename",
                "document_identity",
            ]
            with manifest.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                writer.writerows(
                    [
                        {
                            "source_id": "SRC0015",
                            "family": "veterinary_drug",
                            "standard": "MOA Announcement 235",
                            "version_year": "2002",
                            "source_role": "edition_source",
                            "release_status": "extracted",
                            "source_filename": "农业部235号公告-2002 动物性食品中兽药最高残留限量.doc",
                            "document_identity": "official_announcement",
                        },
                        {
                            "source_id": "SRC0016",
                            "family": "veterinary_drug",
                            "standard": "GB 31650",
                            "version_year": "2019",
                            "source_role": "edition_source",
                            "release_status": "extracted",
                            "source_filename": "GB 31650-2019 食品安全国家标准 食品中兽药最大残留限量.pdf",
                            "document_identity": "full_standard",
                        },
                        {
                            "source_id": "SRC0017",
                            "family": "veterinary_drug",
                            "standard": "GB 31650.1",
                            "version_year": "2022",
                            "source_role": "edition_source",
                            "release_status": "extracted",
                            "source_filename": "GB 31650.1-2022 食品安全国家标准 食品中兽药最大残留限量增补.pdf",
                            "document_identity": "supplementary_standard",
                        },
                    ]
                )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--coverage-manifest",
                    str(manifest),
                    "--output-dir",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            with (output / "standard_versions_long.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                rows = list(csv.DictReader(handle))
            identity = [
                (row["std_id"], row["number"], row["year"], row["event_type"])
                for row in rows
            ]
            self.assertIn(("GB 31650", "31650", "2019", "正式版"), identity)
            self.assertIn(("GB 31650.1", "31650.1", "2022", "正式版"), identity)
            self.assertIn(
                ("MOA Announcement 235", "235", "2002", "正式版"), identity
            )
            self.assertEqual(len(identity), len(set(identity)))

    def test_identity_audit_copies_do_not_duplicate_official_amendments(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "catalog.csv"
            workbook = root / "official.xlsx"
            anchor = root / "anchor"
            manifest = root / "source_manifest.csv"
            output = root / "output"
            category = anchor / "污染物限量"
            category.mkdir(parents=True)
            (category / "GB 2762-2017.pdf").touch()
            (category / "GB 2762-2022.pdf").touch()

            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerows(
                    [
                        {
                            "code": "GB",
                            "number": "2762",
                            "year": "2017",
                            "title": "食品安全国家标准 食品中污染物限量",
                            "path": "anchor/污染物限量/GB 2762-2017.pdf",
                        },
                        {
                            "code": "GB",
                            "number": "2762",
                            "year": "2017",
                            "title": "GB 2762-2017第1号修改单",
                            "path": "anchor/污染物限量/GB 2762-2017第1号修改单.pdf",
                        },
                        {
                            "code": "GB",
                            "number": "2762",
                            "year": "2022",
                            "title": "食品安全国家标准 食品中污染物限量",
                            "path": "anchor/污染物限量/GB 2762-2022.pdf",
                        },
                        {
                            "code": "GB",
                            "number": "2762",
                            "year": "2022",
                            "title": "GB 2762-2022第1号修改单",
                            "path": "anchor/污染物限量/GB 2762-2022第1号修改单.pdf",
                        },
                    ]
                )

            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            sheet.append(
                [
                    "GB 2762-2017",
                    "食品安全国家标准 食品中污染物限量",
                    "2017-03-17",
                    "2017-09-17",
                    "污染物限量",
                ]
            )
            sheet.append(
                [
                    "GB 2762-2017",
                    "《食品安全国家标准 食品中污染物限量》（GB 2762-2017）第1号修改单",
                    "2021-02-22",
                    "2021-02-22",
                    "污染物限量",
                ]
            )
            sheet.append(
                [
                    "GB 2762-2022",
                    "食品安全国家标准 食品中污染物限量",
                    "2022-06-30",
                    "2023-06-30",
                    "污染物限量",
                ]
            )
            sheet.append(
                [
                    "GB 2762-2022",
                    "GB 2762-2022《食品安全国家标准 食品中污染物限量》第1号修改单",
                    "2025-03-16",
                    "2025-03-16",
                    "污染物限量",
                ]
            )
            book.save(workbook)
            book.close()

            fields = [
                "source_id",
                "family",
                "standard",
                "version_year",
                "source_role",
                "release_status",
                "source_filename",
                "document_identity",
            ]
            with manifest.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                writer.writerows(
                    [
                        {
                            "source_id": "SRC0003",
                            "family": "contaminant",
                            "standard": "GB 2762",
                            "version_year": "2017",
                            "source_role": "edition_source",
                            "release_status": "extracted",
                            "source_filename": "GB 2762-2017 食品安全国家标准 食品中污染物限量.pdf",
                            "document_identity": "full_standard",
                        },
                        {
                            "source_id": "SRC0022",
                            "family": "contaminant",
                            "standard": "GB 2762",
                            "version_year": "2017",
                            "source_role": "identity_audit_copy",
                            "release_status": "excluded_from_extraction",
                            "source_filename": "GB 2762-2017第1号修改单.pdf",
                            "document_identity": "full_standard_text_misnamed_as_amendment",
                        },
                        {
                            "source_id": "SRC0004",
                            "family": "contaminant",
                            "standard": "GB 2762",
                            "version_year": "2022",
                            "source_role": "edition_source",
                            "release_status": "extracted",
                            "source_filename": "GB 2762-2022 食品安全国家标准 食品中污染物限量.pdf",
                            "document_identity": "full_standard",
                        },
                        {
                            "source_id": "SRC0023",
                            "family": "contaminant",
                            "standard": "GB 2762",
                            "version_year": "2022",
                            "source_role": "identity_audit_copy",
                            "release_status": "excluded_from_extraction",
                            "source_filename": "GB 2762-2022第1号修改单.pdf",
                            "document_identity": "full_standard_text_misnamed_as_amendment",
                        },
                    ]
                )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--coverage-manifest",
                    str(manifest),
                    "--output-dir",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            with (output / "standard_versions_long.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                rows = list(csv.DictReader(handle))
            target = [row for row in rows if row["std_id"] == "GB 2762"]
            official_versions = [
                row for row in target if row["event_type"] == "正式版"
            ]
            amendments = [row for row in target if row["event_type"] == "修改单"]
            self.assertEqual(len(official_versions), 2)
            self.assertEqual(len(amendments), 2)
            self.assertEqual(
                {(row["year"], row["amend_no"]) for row in amendments},
                {("2017", "第1号"), ("2022", "第1号")},
            )
            self.assertEqual(
                {row["year"]: row["发布日期"] for row in amendments},
                {"2017": "2021-02-22", "2022": "2025-03-16"},
            )

    def test_official_duplicates_merge_or_emit_auditable_conflict(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "catalog.csv"
            workbook = root / "official.xlsx"
            anchor = root / "anchor"
            output = root / "output"
            additive = anchor / "食品添加剂"
            related_product = anchor / "食品相关产品"
            additive.mkdir(parents=True)
            related_product.mkdir(parents=True)
            filenames = (
                "GB 1886.67-2015 食品添加剂 皂荚糖胶.pdf",
                "GB 1886.67-2015 食品安全国家标准 食品添加剂 皂荚糖胶.pdf",
                "GB 14936-2012 食品安全国家标准 食品添加剂硅藻土.pdf",
            )
            for filename in filenames:
                (additive / filename).touch()
            related_filename = "GB 14936-2012 食品安全国家标准 硅藻土.pdf"
            (related_product / related_filename).touch()

            catalog_rows = [
                {
                    "code": "GB",
                    "number": "1886.67",
                    "year": "2015",
                    "title": "食品添加剂 皂荚糖胶",
                    "path": f"anchor/食品添加剂/{filenames[0]}",
                },
                {
                    "code": "GB",
                    "number": "1886.67",
                    "year": "2015",
                    "title": "食品安全国家标准 食品添加剂 皂荚糖胶",
                    "path": f"anchor/食品添加剂/{filenames[1]}",
                },
                {
                    "code": "GB",
                    "number": "14936",
                    "year": "2012",
                    "title": "食品安全国家标准 食品添加剂硅藻土",
                    "path": f"anchor/食品添加剂/{filenames[2]}",
                },
                {
                    "code": "GB",
                    "number": "14936",
                    "year": "2012",
                    "title": "食品安全国家标准  硅藻土",
                    "path": f"anchor/食品相关产品/{related_filename}",
                },
            ]
            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerows(catalog_rows)

            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            sheet.append(
                [
                    "GB 1886.67-2015",
                    "食品添加剂 皂荚糖胶",
                    "2015-09-22",
                    "2016-03-22",
                    "食品添加剂",
                ]
            )
            sheet.append(
                [
                    "GB 1886.67-2015",
                    "食品安全国家标准 食品添加剂 皂荚糖胶",
                    "2015-09-22",
                    "2016-03-22",
                    "食品添加剂",
                ]
            )
            sheet.append(
                [
                    "GB 14936-2012",
                    "食品安全国家标准 食品添加剂硅藻土",
                    "2012-12-25",
                    "2013-01-25",
                    "食品添加剂",
                ]
            )
            sheet.append(
                [
                    "GB 14936-2012",
                    "食品安全国家标准  硅藻土",
                    "2012-05-17",
                    "2012-11-17",
                    "食品相关产品",
                ]
            )
            book.save(workbook)
            book.close()

            command = [
                sys.executable,
                str(SCRIPT),
                "--catalog-csv",
                str(catalog),
                "--official-xlsx",
                str(workbook),
                "--anchor-dir",
                str(anchor),
                "--output-dir",
                str(output),
            ]
            result = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            with (output / "standard_versions_long.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                events = list(csv.DictReader(handle))
            with (output / "standard_timeline_conflicts.csv").open(
                encoding="utf-8-sig", newline=""
            ) as handle:
                conflicts = list(csv.DictReader(handle))

            equivalent = [row for row in events if row["std_id"] == "GB 1886.67"]
            selected = [row for row in events if row["std_id"] == "GB 14936"]
            self.assertEqual(len(equivalent), 1)
            self.assertEqual(len(selected), 1)
            self.assertEqual(selected[0]["发布日期"], "2012-12-25")
            self.assertFalse(
                any(row["std_id"] == "GB 1886.67" for row in conflicts)
            )
            target_conflicts = [
                row for row in conflicts if row["std_id"] == "GB 14936"
            ]
            self.assertEqual(len(target_conflicts), 2)
            self.assertEqual(
                sum(row["selected"] == "1" for row in target_conflicts),
                1,
            )
            rejected = [row for row in target_conflicts if row["selected"] == "0"]
            self.assertEqual(rejected[0]["official_category"], "食品相关产品")
            self.assertIn("stable_rank", rejected[0]["selection_reason"])

            second_output = root / "second-output"
            second_command = [
                str(second_output) if value == str(output) else value for value in command
            ]
            second_result = subprocess.run(
                second_command,
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(second_result.returncode, 0, second_result.stderr)
            for filename in (
                "standard_versions_long.csv",
                "standard_timeline.csv",
                "standard_timeline_conflicts.csv",
            ):
                self.assertNotIn(b"\r\n", (output / filename).read_bytes())
                self.assertEqual(
                    (output / filename).read_bytes(),
                    (second_output / filename).read_bytes(),
                )

    def test_duplicate_manifest_edition_identity_fails_closed(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            catalog = root / "catalog.csv"
            workbook = root / "official.xlsx"
            anchor = root / "anchor"
            manifest = root / "source_manifest.csv"
            output = root / "output"
            category = anchor / "污染物限量"
            category.mkdir(parents=True)
            (category / "GB 2762-2022.pdf").touch()
            with catalog.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(
                    handle,
                    fieldnames=["code", "number", "year", "title", "path"],
                )
                writer.writeheader()
                writer.writerow(
                    {
                        "code": "GB",
                        "number": "2762",
                        "year": "2022",
                        "title": "食品安全国家标准 食品中污染物限量",
                        "path": "anchor/污染物限量/GB 2762-2022.pdf",
                    }
                )
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "全量汇总"
            sheet.append(["标准号", "名称", "发布日期", "实施日期", "分类"])
            book.save(workbook)
            book.close()
            fields = [
                "source_id",
                "family",
                "standard",
                "version_year",
                "source_role",
                "release_status",
                "source_filename",
                "document_identity",
            ]
            with manifest.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                for source_id, filename in (
                    ("SRC0001", "GB 31650-2019-a.pdf"),
                    ("SRC0002", "GB 31650-2019-b.pdf"),
                ):
                    writer.writerow(
                        {
                            "source_id": source_id,
                            "family": "veterinary_drug",
                            "standard": "GB 31650",
                            "version_year": "2019",
                            "source_role": "edition_source",
                            "release_status": "extracted",
                            "source_filename": filename,
                            "document_identity": "full_standard",
                        }
                    )

            result = subprocess.run(
                [
                    sys.executable,
                    str(SCRIPT),
                    "--catalog-csv",
                    str(catalog),
                    "--official-xlsx",
                    str(workbook),
                    "--anchor-dir",
                    str(anchor),
                    "--coverage-manifest",
                    str(manifest),
                    "--output-dir",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("duplicate edition_source identity", result.stderr)


if __name__ == "__main__":
    unittest.main()
