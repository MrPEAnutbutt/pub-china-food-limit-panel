#!/usr/bin/env python3
"""Rebuild the 2002 veterinary-drug MRL table from the official Word file."""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import re
import shutil
import tempfile

from legacy_family_replay import (
    ACQUISITION_FIELDS,
    CSVSpec,
    file_metadata,
    load_locked_csv,
    python_environment,
    read_manifest,
    ReplayError,
    run_checked,
    tool_version,
    validate_source_acquisition,
    verify_manifest_source,
    write_csv_atomic,
    write_json_atomic,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "Data" / "Interim" / "vetdrug_rebuild"
DEFAULT_MANIFEST = (
    PROJECT_ROOT / "Data" / "Final" / "china_food_limit_panel_v1" / "source_manifest.csv"
)
LONG_FIELDS = (
    "year",
    "drug",
    "english",
    "animal_raw",
    "animal",
    "tissue",
    "limit_value",
    "unit",
    "source",
)
QUARANTINE_FIELDS = (
    "drug",
    "english",
    "reject",
    "animal_paras",
    "tissue_paras",
    "value_paras",
)
LOCATOR_FIELDS = (
    "output_row",
    "source_table_index",
    "source_table_row",
    "animal_block_index",
    "tissue_index",
    "drug",
    "english",
    "animal_raw",
    "tissue",
    "limit_value",
)
SOURCE_ACQUISITION_SPEC = CSVSpec(
    fields=ACQUISITION_FIELDS,
    rows=3,
    unique_key=("family", "version_year"),
    sha256="a5e35c1f90b3995cc8fa3f18ddb7fe0377dc21cf868a40b99aeb229f58277ffc",
)
EXPECTED_OUTPUTS = {
    "vetdrug_2002_long.csv": (
        551,
        "eb846756b187405a430425d25954f7dd68e708f8944d7e82a2aad7ecd142fa78",
    ),
    "vetdrug_2002_quarantine.csv": (
        0,
        "b40e78bb867c90976d324d7f6971ca59c86ab36939af1f005f7c337b638bca6b",
    ),
    "vetdrug_2002_row_locators.csv": (
        551,
        "bc5e299f6320b124b3ff5e5ea6ae594246d22a4faad32056719a0f0c5e179aea",
    ),
}
TISSUE_FIX = {"＋": "+", "：": ":", "／": "/"}
NUMERIC = re.compile(r"^\d+(?:\.\d+)?$")
ANIMAL_VOCAB = {
    "牛",
    "羊",
    "绵羊",
    "山羊",
    "猪",
    "马",
    "兔",
    "鹿",
    "禽",
    "家禽",
    "鸡",
    "火鸡",
    "鸭",
    "鹅",
    "鱼",
    "虾",
    "鲑鱼",
    "鳟鱼",
    "蜜蜂",
    "所有食品动物",
    "其他动物",
    "所有动物",
    "其他",
}
ANIMAL_ALIAS = {"其他": "其他动物", "所有动物": "所有食品动物"}


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-doc", type=Path, required=True)
    parser.add_argument("--source-manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument(
        "--source-acquisition",
        type=Path,
        default=DEFAULT_OUTPUT_DIR / "source_acquisition.csv",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--soffice", default="soffice")
    return parser.parse_args(argv)


def normalize_tissue(value: str) -> str:
    for source, replacement in TISSUE_FIX.items():
        value = value.replace(source, replacement)
    value = re.sub(r"\s+", "", value)
    return {
        "可食性组织": "可食组织",
        "其他可食组织": "可食组织",
        "脂": "脂肪",
    }.get(value, value)


def split_blocks(paragraphs: list[str]) -> list[list[str]]:
    output: list[list[str]] = []
    current: list[str] = []
    for paragraph in paragraphs:
        if paragraph:
            current.append(paragraph)
        elif current:
            output.append(current)
            current = []
    if current:
        output.append(current)
    return output


def cell_paragraphs(cell: object) -> list[str]:
    return [re.sub(r"\s+", "", paragraph.text) for paragraph in cell.paragraphs]


def complete_animal(value: str) -> bool:
    tokens = [
        token
        for token in re.sub(r"[（(].*?[)）]", "", value).split("/")
        if token
    ]
    return bool(tokens) and all(token in ANIMAL_VOCAB for token in tokens)


def normalize_animal(value: str) -> str:
    return "/".join(ANIMAL_ALIAS.get(token, token) for token in value.split("/") if token)


def animal_entries(paragraphs: list[str]) -> list[str]:
    entries: list[str] = []
    for paragraph in [value for value in paragraphs if value]:
        if entries and (paragraph.startswith("/") or not complete_animal(entries[-1])):
            entries[-1] += paragraph
        else:
            entries.append(paragraph)
    return entries


def convert_doc_to_docx(
    source_doc: Path,
    temporary_dir: Path,
    *,
    soffice: str,
) -> Path:
    """Convert a copied source document with an isolated LibreOffice profile."""

    copied_doc = temporary_dir / "src.doc"
    shutil.copy2(source_doc, copied_doc)
    profile = temporary_dir / "libreoffice-profile"
    profile.mkdir()
    run_checked(
        [
            soffice,
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            "docx",
            copied_doc,
            "--outdir",
            temporary_dir,
        ],
        timeout=180,
    )
    output = temporary_dir / "src.docx"
    if not output.is_file():
        raise ReplayError("LibreOffice returned success without a DOCX output")
    return output


def extract_rows(
    document: object,
) -> tuple[
    list[dict[str, str]],
    list[dict[str, str]],
    list[dict[str, str]],
]:
    tables = [table for table in document.tables if len(table.columns) == 5]
    if len(tables) != 1:
        raise ReplayError(
            f"expected one five-column Appendix 2 table, found {len(tables)}"
        )
    table = tables[0]
    accepted: list[dict[str, str]] = []
    quarantined: list[dict[str, str]] = []
    locators: list[dict[str, str]] = []
    for source_table_row, row in enumerate(table.rows[1:], start=2):
        cells = row.cells
        first = [paragraph for paragraph in cell_paragraphs(cells[0]) if paragraph]
        if not first:
            continue
        chinese = next(
            (paragraph for paragraph in first if re.match(r"^[一-鿿]", paragraph)),
            "",
        )
        english = next(
            (
                paragraph
                for paragraph in first
                if re.match(r"^[A-Za-z]", paragraph)
                and not paragraph.upper().startswith("ADI")
            ),
            "",
        )
        drug = chinese or english
        if not drug or drug.startswith("药物名"):
            continue

        animals = animal_entries(cell_paragraphs(cells[2]))
        tissue_blocks = split_blocks(cell_paragraphs(cells[3]))
        value_blocks = split_blocks(cell_paragraphs(cells[4]))
        if len(animals) == 1 and len(tissue_blocks) > 1:
            animals *= len(tissue_blocks)

        rejection = ""
        if not animals or not tissue_blocks or not value_blocks:
            rejection = "动物/组织/限量有空列"
        elif not (len(animals) == len(tissue_blocks) == len(value_blocks)):
            rejection = (
                f"块数不等(动物{len(animals)}/组织{len(tissue_blocks)}/"
                f"限量{len(value_blocks)})"
            )
        elif any(
            len(tissues) != len(values)
            for tissues, values in zip(tissue_blocks, value_blocks)
        ):
            rejection = "块内组织数≠限量数"
        elif not all(complete_animal(animal) for animal in animals):
            rejection = "动物名不在闭合词表内(疑折行或新动物)"
        if rejection:
            quarantined.append(
                {
                    "drug": drug,
                    "english": english,
                    "reject": rejection,
                    "animal_paras": "|".join(animals),
                    "tissue_paras": "|".join(
                        "/".join(block) for block in tissue_blocks
                    ),
                    "value_paras": "|".join(
                        "/".join(block) for block in value_blocks
                    ),
                }
            )
            continue

        for animal_block_index, (animal_raw, tissues, values) in enumerate(
            zip(animals, tissue_blocks, value_blocks), start=1
        ):
            animal = normalize_animal(
                re.sub(r"[（(].*?[)）]", "", animal_raw).strip()
            )
            for tissue_index, (tissue, value) in enumerate(
                zip(tissues, values), start=1
            ):
                if not NUMERIC.match(value):
                    continue
                normalized_tissue = normalize_tissue(tissue)
                accepted.append(
                    {
                        "year": "2002",
                        "drug": drug,
                        "english": english.lower(),
                        "animal_raw": animal_raw,
                        "animal": animal,
                        "tissue": normalized_tissue,
                        "limit_value": value,
                        "unit": "μg/kg",
                        "source": "docx-block",
                    }
                )
                locators.append(
                    {
                        "output_row": str(len(accepted)),
                        "source_table_index": "1",
                        "source_table_row": str(source_table_row),
                        "animal_block_index": str(animal_block_index),
                        "tissue_index": str(tissue_index),
                        "drug": drug,
                        "english": english.lower(),
                        "animal_raw": animal_raw,
                        "tissue": normalized_tissue,
                        "limit_value": value,
                    }
                )
    return accepted, quarantined, locators


def main(argv: list[str] | None = None) -> None:
    args = parse_args(argv)
    manifest_rows = read_manifest(args.source_manifest)
    source_metadata = verify_manifest_source(
        args.source_doc,
        manifest_rows,
        family="veterinary_drug",
        version_year="2002",
    )
    acquisition_rows = load_locked_csv(
        args.source_acquisition, SOURCE_ACQUISITION_SPEC
    )
    acquisition_2002 = [
        row
        for row in acquisition_rows
        if row["family"] == "veterinary_drug" and row["version_year"] == "2002"
    ]
    validate_source_acquisition(acquisition_2002, [source_metadata])

    import docx

    with tempfile.TemporaryDirectory() as directory:
        docx_path = convert_doc_to_docx(
            args.source_doc, Path(directory), soffice=args.soffice
        )
        document = docx.Document(docx_path)
    accepted, quarantined, locators = extract_rows(document)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    long_name = "vetdrug_2002_long.csv"
    expected_rows, expected_digest = EXPECTED_OUTPUTS[long_name]
    long_metadata = write_csv_atomic(
        args.output_dir / long_name,
        LONG_FIELDS,
        accepted,
        expected_rows=expected_rows,
        expected_sha256=expected_digest,
    )
    quarantine_name = "vetdrug_2002_quarantine.csv"
    expected_rows, expected_digest = EXPECTED_OUTPUTS[quarantine_name]
    quarantine_metadata = write_csv_atomic(
        args.output_dir / quarantine_name,
        QUARANTINE_FIELDS,
        quarantined,
        expected_rows=expected_rows,
        expected_sha256=expected_digest,
    )
    locator_name = "vetdrug_2002_row_locators.csv"
    expected_rows, expected_digest = EXPECTED_OUTPUTS[locator_name]
    locator_metadata = write_csv_atomic(
        args.output_dir / locator_name,
        LOCATOR_FIELDS,
        locators,
        expected_rows=expected_rows,
        expected_sha256=expected_digest,
    )

    report = {
        "schema_version": 1,
        "pipeline": "vetdrug_2002_word_rebuild",
        "status": "pass",
        "environment": {
            **python_environment(),
            "python_docx": docx.__version__,
            "libreoffice": tool_version([args.soffice, "--version"]),
        },
        "source_manifest": file_metadata(args.source_manifest),
        "sources": [source_metadata],
        "source_acquisition": {
            **file_metadata(args.source_acquisition),
            "rows": len(acquisition_rows),
            "matched_rows": len(acquisition_2002),
            "identity_verified": True,
        },
        "conversion": {
            "filename": "src.docx",
            "retained": False,
            "digest_gate": False,
            "reason": (
                "Temporary DOCX container metadata may vary; the reconstructed "
                "CSV identities are the deterministic conversion gates."
            ),
        },
        "checks": {
            "accepted_rows": len(accepted),
            "unique_drugs": len({row["drug"] for row in accepted}),
            "blank_animal_rows": sum(not row["animal"] for row in accepted),
            "quarantined_drugs": len(quarantined),
            "animal_distribution": dict(
                Counter(row["animal"] for row in accepted).most_common()
            ),
            "source_identity_verified": True,
            "output_identity_verified": True,
        },
        "outputs": [long_metadata, quarantine_metadata, locator_metadata],
        "replay_scope": {
            "source_table_reextraction": True,
            "downstream_rebuild": True,
            "manual_seed_replay": False,
            "independent_retranscription": False,
            "description": (
                "The replay converts the byte-identified official Word file, "
                "re-extracts its five-column table, and rebuilds the frozen 2002 "
                "CSV plus row-locator sidecar. It does not constitute an "
                "independent second transcription."
            ),
        },
        "manual_transcription_gap": {
            "status": "none",
            "description": (
                "No manual transcription seed contributes to the 2002 output; "
                "rows are reconstructed from the Word table stream."
            ),
        },
        "row_locator_gap": {
            "status": "open",
            "records": len(locators),
            "available_fields": [
                "source_table_index",
                "source_table_row",
                "animal_block_index",
                "tissue_index",
            ],
            "missing_fields": ["source_page", "printed_row_label"],
            "description": (
                "The sidecar identifies the Word table row and within-cell block "
                "for every output row. The binary DOC has no stable rendered-page "
                "locator, and the printed table does not label each row."
            ),
        },
    }
    write_json_atomic(args.output_dir / "vetdrug_2002_replay_report.json", report)
    print(
        f"2002 veterinary-drug replay passed: {len(accepted)} rows, "
        f"{len(quarantined)} quarantined drugs"
    )


if __name__ == "__main__":
    main()
