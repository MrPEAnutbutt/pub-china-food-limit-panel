#!/usr/bin/env python3
"""Build the independent reviewer's handoff folder and unpack what comes back.

The folder holds two value-free workbooks, the source PDFs, a Chinese guide and
a declaration form. Part 1 is the 144-row blind re-entry packet
(`Data/Interim/manual_reentry/blind_packet`, protocol 1.2); part 2 is the
300-row reading of the GB 2763-2019 and -2021 scans
(`Data/Interim/pesticide_audit`). The folder contains source PDFs and stays
private; it is never committed or deposited.

Usage:
    python3 Program/Clean/261024_reviewer_handoff.py build --source-root Data/Raw
    python3 Program/Clean/261024_reviewer_handoff.py unpack \\
        --part1 <returned 01 workbook> --part2 <returned 02 workbook> \\
        --output-dir <empty-dir>

Round 2 holds part 2 only, drawn again in `Data/Interim/pesticide_audit_round2`:
    python3 Program/Clean/261024_reviewer_handoff.py build --source-root Data/Raw \\
        --part2-only --audit-dir Data/Interim/pesticide_audit_round2 \\
        --output-dir Data/Interim/independent_review_round2/核验包
    python3 Program/Clean/261024_reviewer_handoff.py unpack \\
        --audit-dir Data/Interim/pesticide_audit_round2 \\
        --part2 <returned workbook> --output-dir <empty-dir>

`unpack` records the digest of each returned workbook before reading it,
rejects any change to a locked locator cell, and writes the completed CSVs in
the template format. With `--check`, it only reports format problems (used on
the reviewer's first ten rows; it never compares values).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_ROOT = Path(__file__).resolve().parents[2]
BLIND_PACKET = PROJECT_ROOT / "Data" / "Interim" / "manual_reentry" / "blind_packet"
AUDIT_DIR = PROJECT_ROOT / "Data" / "Interim" / "pesticide_audit"
MANIFEST = (
    PROJECT_ROOT / "Data" / "Final" / "china_food_limit_panel_v1" / "source_manifest.csv"
)
DEFAULT_OUTPUT = PROJECT_ROOT / "Data" / "Interim" / "independent_review" / "核验包"

PART1_FILE = "01_第一部分_手录144行_录入表.xlsx"
PART2_FILE = "02_第二部分_农残300行_盲读表.xlsx"
ROUND2_FILE = "01_农残第二轮{rows}行_盲读表.xlsx"
PART1_SHEETS = (
    ("真菌毒素2005", "mycotoxin_2005_blind_entry.csv", "SRC0012"),
    ("兽药2019", "vetdrug_2019_blind_entry.csv", "SRC0016"),
    ("兽药2022", "vetdrug_2022_blind_entry.csv", "SRC0017"),
)
PART2_SHEET = ("农残", "pesticide_blind_read_template.csv")
PART2_SOURCES = ("SRC0009", "SRC0010")
LOCATORS = {
    "blind_record_id", "source_order", "source_id", "section_number", "pdf_page",
    "printed_page", "printed_table", "printed_row", "continued_from_previous_page",
    "blind_id", "edition", "pesticide_zh", "parent_category_zh", "food_zh",
    "reference_scan_pdf_page",
}
LABELS = {
    "blind_record_id": "编号", "source_order": "顺序", "source_id": "源文件",
    "section_number": "条款号", "pdf_page": "PDF页", "printed_page": "印刷页码",
    "printed_table": "表号", "printed_row": "表内第几行",
    "continued_from_previous_page": "接上页", "entry_status": "状态",
    "toxin_raw": "真菌毒素", "food_category_raw": "食品类别(名称)",
    "limit_value_raw": "限量值", "unit_raw": "单位", "drug_zh_raw": "药物中文名",
    "drug_en_raw": "药物英文名", "animal_raw": "动物种类", "tissue_raw": "靶组织",
    "reviewer_note": "备注", "blind_id": "编号", "edition": "年版",
    "pesticide_zh": "农药", "parent_category_zh": "食品类别", "food_zh": "食品名称",
    "reference_scan_pdf_page": "参考PDF页",
}
STATUS_PART1 = {"已录入": "transcribed", "看不清": "illegible"}
STATUS_PART2 = {"已录入": "transcribed", "看不清": "illegible", "找不到": "not_found"}
GREY = PatternFill("solid", fgColor="E7E6E6")

GUIDE = """\
# 数据核验说明

感谢帮忙。这是一份中国食品安全国家标准限量数据库的独立核验：请你对着标准原件，把指定的行录下来。表里只给每一行在原件里的位置，不给建库时录的数值；你交回以后，作者再把两边逐格比对，不一致的地方回到原件裁决。你录的结果会以匿名方式写进论文和数据说明。

## 包里有什么

- `01_第一部分_手录144行_录入表.xlsx`：三个工作表，真菌毒素 2005 年 9 行、兽药 2019 年 13 行、兽药 2022 年 122 行。
- `02_第二部分_农残300行_盲读表.xlsx`：农药残留限量 2019 年、2021 年两版，共 300 行。
- `源文件/`：第一部分 3 份标准、第二部分 2 份扫描件。只用这里的文件，不要另找版本。
- `03_核验声明.docx`：做完后填写身份、日期并逐条勾选，不用写姓名。

## 通用规则

1. 只填白色格子。灰色的是定位信息，已锁定，不能改。
2. 照原文录，不换算、不纠错、不补全：数字照印刷写法（印的是 0.10 就写 0.10）；带 \\* 号的连 \\* 号一起录；数值后面印有括号说明的也录上。
3. 页码指 PDF 阅读器里的页码（第几页 / 共几页），不是页面底部印的页码。
4. 看不清就在「状态」选「看不清」，在「备注」写原因，其他格留空，不要猜。2019 年农残扫描件有大块水印，遇到被盖住的数字照此处理。
5. 合并单元格（一个类别或动物跨好几行）：每一行都把它继承的名称照录。
6. 录入期间保持独立：不看这个数据库的发布文件、代码仓库和论文稿；不和作者讨论具体数值；不用 OCR、AI 等识别工具，每一格都由你本人看原件录入。格式上的疑问可以问作者。

## 第一部分怎么录

每行已给出条款号、PDF 页、表号和表内第几行。表内行号从该表第一行数据数起，表格跨页时接着往下数；「接上页」为 true 的行在下一页的续表里。按行录：

- 真菌毒素表：真菌毒素名称、食品类别（名称）、限量值、单位，照原文写法。
- 兽药表：药物中文名、英文名（括号里的英文，印了几个就录几个）、动物种类、靶组织、限量值、单位。同一药物下的每一行都要录药名。

## 第二部分怎么录

每行给出年版、条款号（4.后面的数字）、农药、食品类别、食品名称和参考 PDF 页（实际位置可能前后差一页）。在该条款的限量表里找到这个食品，录：

- 限量值：照印刷写法，带 \\* 号的连 \\* 号一起录。
- 状态：找到并录好选「已录入」；数字看不清选「看不清」；这一条款里找不到这个食品（或类别不对）选「找不到」，并在备注写你查了哪几页。
- 印刷的食品或类别写法和表里不一样时，在备注照录印刷写法。

## 怎么交

1. 先试录：两部分各录前 10 行，发给作者。作者只看格式和理解，不告诉对错。
2. 没问题再录完。全部录完后，把两份 xlsx 和填好的声明一起发给作者。交回后不再修改。
"""

DECLARATION = """\
# 核验声明 Verification Declaration

**核验人匿名。**本声明不记录核验人的姓名、签名或联系方式，随录入表一起发回作者。论文、数据说明和存档材料中只用化名编号指代核验人。\\
**The verifier is anonymous.** This declaration records no name, signature or contact details and is returned to the author together with the entry forms. The manuscript, the data documentation and all archived records refer to the verifier only by a pseudonymous code.

单位 / Affiliation: \\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_

完成日期 / Date of completion: \\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_

请逐条确认（在 □ 处打 √）：\\
Please confirm each statement by ticking (√) its box (□):

□ 1. 我没有参与这个数据库的建库工作（录入、抽取、清洗或核对）。\\
I took no part in building this database (data entry, extraction, cleaning or checking).

□ 2. 在交回录入表之前，我没有看过建库时的录入值、数据库的发布文件、代码仓库或任何比对结果。\\
Before returning the entry forms, I did not see the values entered when the database was built, the released data files, the code repository or any comparison result.

□ 3. 我只使用随包提供的源文件。\\
I used only the source files supplied with the package.

□ 4. 两部分的每一行都由我本人对着源文件独立录入，没有借助他人，也没有使用 OCR、AI 等识别工具。\\
I entered every row of both parts myself, independently and from the source files, without help from others and without OCR, AI or other recognition tools.

□ 5. 交回后我不再修改录入表。\\
I will not change the entry forms after returning them.

□ 6. 我同意作者在论文和数据说明中以匿名方式描述我的核验工作。\\
I agree that the author may describe my verification work anonymously in the manuscript and the data documentation.

第一部分确认：在录好的录入表定版之前，我没有看过建库时的录入值和数据库发布的数值，指定的源文件行都由我独立录入。\\
Confirmation for part 1: I independently transcribed the specified source rows without viewing the construction seeds or released values before the completed entry files were frozen.
"""

GUIDE_ROUND2 = """\
# 数据核验说明（第二轮）

谢谢你完成第一轮。第一轮的比对发现了一类建库错误，作者修正了抽取规则，并按事先定好的规程，从同样几部分里再抽一批你没录过的行，请你再录一轮。

## 包里有什么

- `{workbook}`：农药残留限量 2019 年、2021 年两版，共 {rows} 行，都是第一轮没有出现过的行。
- `源文件/`：两份扫描件，和第一轮第二部分的一样。只用这里的文件，不要另找版本。
- `02_核验声明.docx`：做完后填写身份、日期并逐条勾选，不用写姓名。

## 怎么录

规则和第一轮第二部分相同：

1. 只填白色格子。灰色的是定位信息，已锁定，不能改。
2. 照原文录，不换算、不纠错、不补全：数字照印刷写法（印的是 0.10 就写 0.10）；带 \\* 号的连 \\* 号一起录；数值后面印有括号说明的也录上。
3. 页码指 PDF 阅读器里的页码（第几页 / 共几页），不是页面底部印的页码。参考 PDF 页的实际位置可能前后差一页。
4. 看不清就在「状态」选「看不清」，在「备注」写原因，其他格留空，不要猜。2019 年扫描件有大块水印，遇到被盖住的数字照此处理。
5. 在给出的条款的限量表里找到这个食品，录限量值。找到并录好选「已录入」；这一条款里找不到这个食品，或它不在给出的食品类别下，选「找不到」，并在备注写你查了哪几页、原文怎么印的。
6. 印刷的食品或类别写法和表里不一样时，在备注照录印刷写法。
7. 录入期间保持独立：不看这个数据库的发布文件、代码仓库和论文稿；不和作者讨论具体数值；不用 OCR、AI 等识别工具，每一格都由你本人看原件录入。格式上的疑问可以问作者。

## 怎么交

格式和第一轮相同，不用试录。全部录完后，把 xlsx 和填好的声明一起发给作者。交回后不再修改。
"""

DECLARATION_ROUND2 = """\
# 核验声明 Verification Declaration

**核验人匿名。**本声明不记录核验人的姓名、签名或联系方式，随录入表一起发回作者。论文、数据说明和存档材料中只用化名编号指代核验人。\\
**The verifier is anonymous.** This declaration records no name, signature or contact details and is returned to the author together with the entry form. The manuscript, the data documentation and all archived records refer to the verifier only by a pseudonymous code.

单位 / Affiliation: \\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_

完成日期 / Date of completion: \\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_\\_

请逐条确认（在 □ 处打 √）：\\
Please confirm each statement by ticking (√) its box (□):

□ 1. 我没有参与这个数据库的建库工作（录入、抽取、清洗或核对）。\\
I took no part in building this database (data entry, extraction, cleaning or checking).

□ 2. 在交回本轮录入表之前，我没有看过建库时的录入值、数据库的发布文件或代码仓库，也没有看过本轮任何一行的比对结果。\\
Before returning the entry form for this round, I did not see the values entered when the database was built, the released data files or the code repository, nor the comparison result for any row in this round.

□ 3. 我只使用随包提供的源文件。\\
I used only the source files supplied with the package.

□ 4. 每一行都由我本人对着源文件独立录入，没有借助他人，也没有使用 OCR、AI 等识别工具。\\
I entered every row myself, independently and from the source files, without help from others and without OCR, AI or other recognition tools.

□ 5. 交回后我不再修改录入表。\\
I will not change the entry form after returning it.

□ 6. 我同意作者在论文和数据说明中以匿名方式描述我的核验工作。\\
I agree that the author may describe my verification work anonymously in the manuscript and the data documentation.
"""


class HandoffError(RuntimeError):
    """Raised when a handoff input or a returned workbook breaks the contract."""


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        return list(reader.fieldnames or ()), list(reader)


def header(field: str) -> str:
    return f"{LABELS[field]} {field}"


def add_sheet(book: Workbook, title: str, fields: list[str], rows, statuses) -> None:
    sheet = book.create_sheet(title)
    sheet.append([header(field) for field in fields])
    for row in rows:
        sheet.append([row[field] for field in fields])
    for column, field in enumerate(fields, start=1):
        letter = sheet.cell(row=1, column=column).column_letter
        sheet.column_dimensions[letter].width = 22 if field in LOCATORS else 18
        if field == "reviewer_note":
            sheet.column_dimensions[letter].width = 36
        for row_index in range(1, len(rows) + 2):
            cell = sheet.cell(row=row_index, column=column)
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index == 1:
                cell.font = Font(bold=True)
                cell.fill = GREY
            elif field in LOCATORS:
                cell.fill = GREY
            else:
                cell.protection = Protection(locked=False)
    status_column = fields.index("entry_status") + 1
    letter = sheet.cell(row=1, column=status_column).column_letter
    validation = DataValidation(
        type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=True
    )
    validation.add(f"{letter}2:{letter}{len(rows) + 1}")
    sheet.add_data_validation(validation)
    sheet.freeze_panes = sheet.cell(row=2, column=status_column)
    sheet.protection.sheet = True
    sheet.protection.formatColumns = False
    sheet.protection.formatRows = False


def write_workbook(path: Path, sheets, statuses) -> None:
    book = Workbook()
    book.remove(book.active)
    for title, fields, rows in sheets:
        add_sheet(book, title, fields, rows, statuses)
    book.save(path)


def write_docx(path: Path, markdown: str) -> None:
    with tempfile.TemporaryDirectory() as folder:
        source = Path(folder) / "source.md"
        source.write_text(markdown, encoding="utf-8")
        subprocess.run(["pandoc", str(source), "-o", str(path)], check=True)


def checked_sources(source_root: Path, ids) -> dict[str, Path]:
    _, manifest = read_csv(MANIFEST)
    by_id = {row["source_id"]: row for row in manifest}
    found = {}
    for source_id in ids:
        row = by_id[source_id]
        matches = list(source_root.rglob(row["source_filename"]))
        verified = [path for path in matches if sha256(path) == row["sha256"]]
        if not verified:
            raise HandoffError(f"{source_id} not found with the manifest digest")
        found[source_id] = verified[0]
    return found


def build(source_root: Path, output: Path, audit_dir: Path = AUDIT_DIR,
          part2_only: bool = False) -> None:
    if output.exists() and any(output.iterdir()):
        raise HandoffError(f"output directory must be empty: {output}")
    receipt: dict[str, object] = {}
    if not part2_only:
        packet = json.loads((BLIND_PACKET / "packet_manifest.json").read_text())
        if packet["protocol_version"] != "1.2":
            raise HandoffError("blind packet is not protocol 1.2")
        for name, meta in packet["files"].items():
            if name.endswith("_blind_entry.csv") and sha256(BLIND_PACKET / name) != meta["sha256"]:
                raise HandoffError(f"blind template differs from its manifest: {name}")
        receipt["blind_packet_id"] = packet["packet_id"]
    audit = json.loads((audit_dir / "audit_manifest.json").read_text())
    template = audit_dir / PART2_SHEET[1]
    if sha256(template) != audit["outputs"]["template"]["sha256"]:
        raise HandoffError("pesticide template differs from its manifest")
    receipt["pesticide_audit_seed"] = audit["seed"]
    if part2_only:
        receipt["pesticide_audit_manifest_sha256"] = sha256(audit_dir / "audit_manifest.json")
    part1_ids = [] if part2_only else [item[2] for item in PART1_SHEETS]
    sources = checked_sources(source_root, part1_ids + list(PART2_SOURCES))

    output.mkdir(parents=True, exist_ok=True)
    fields, rows = read_csv(template)
    if part2_only:
        workbook = ROUND2_FILE.format(rows=len(rows))
        write_workbook(output / workbook, [(PART2_SHEET[0], fields, rows)], STATUS_PART2)
        write_docx(output / "00_核验说明.docx",
                   GUIDE_ROUND2.format(workbook=workbook, rows=len(rows)))
        write_docx(output / "02_核验声明.docx", DECLARATION_ROUND2)
        folders = (("", PART2_SOURCES),)
    else:
        sheets = []
        for title, filename, _ in PART1_SHEETS:
            part1_fields, part1_rows = read_csv(BLIND_PACKET / filename)
            sheets.append((title, part1_fields, part1_rows))
        write_workbook(output / PART1_FILE, sheets, STATUS_PART1)
        write_workbook(output / PART2_FILE, [(PART2_SHEET[0], fields, rows)], STATUS_PART2)
        write_docx(output / "00_核验说明.docx", GUIDE)
        write_docx(output / "03_核验声明.docx", DECLARATION)
        folders = (("第一部分", part1_ids), ("第二部分", PART2_SOURCES))
    for part, ids in folders:
        folder = output / "源文件" / part
        folder.mkdir(parents=True, exist_ok=True)
        for source_id in ids:
            shutil.copy2(sources[source_id], folder / sources[source_id].name)
    receipt["files"] = {
        str(path.relative_to(output)): sha256(path)
        for path in sorted(output.rglob("*")) if path.is_file()
    }
    (output.parent / "handoff_receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_sheet(book, title, template_path: Path, statuses) -> tuple[list[str], list, list]:
    fields, expected = read_csv(template_path)
    if title not in book.sheetnames:
        raise HandoffError(f"sheet missing: {title}")
    values = list(book[title].iter_rows(values_only=True))
    if [cell_text(value) for value in values[0][: len(fields)]] != [header(f) for f in fields]:
        raise HandoffError(f"{title}: header changed")
    rows = [dict(zip(fields, map(cell_text, row))) for row in values[1:]]
    rows = [row for row in rows if any(row.values())]
    if len(rows) != len(expected):
        raise HandoffError(f"{title}: {len(rows)} rows, expected {len(expected)}")
    problems = []
    for row, locator in zip(rows, expected):
        changed = [f for f in fields if f in LOCATORS and row[f] != locator[f]]
        if changed:
            raise HandoffError(f"{title}: locator changed in {locator[fields[0]]}: {changed}")
        status = row["entry_status"]
        if status and status not in statuses:
            problems.append(f"{title} {row[fields[0]]}: unknown status {status!r}")
        row["entry_status"] = statuses.get(status, status)
    return fields, rows, problems


def unpack(part1: Path | None, part2: Path | None, output: Path | None, check: bool,
           audit_dir: Path = AUDIT_DIR) -> int:
    receipt = {"workbooks": {}, "csv": {}}
    jobs = []
    if part1:
        receipt["workbooks"][part1.name] = sha256(part1)
        book = load_workbook(part1)
        for title, filename, _ in PART1_SHEETS:
            jobs.append((filename, *read_sheet(book, title, BLIND_PACKET / filename,
                                               STATUS_PART1)))
    if part2:
        receipt["workbooks"][part2.name] = sha256(part2)
        book = load_workbook(part2)
        jobs.append(("pesticide_blind_read_completed.csv",
                     *read_sheet(book, PART2_SHEET[0], audit_dir / PART2_SHEET[1],
                                 STATUS_PART2)))
    problems = [problem for job in jobs for problem in job[3]]
    for filename, fields, rows, _ in jobs:
        filled = [row for row in rows if row["entry_status"]]
        print(f"{filename}: {len(filled)}/{len(rows)} rows have a status")
    for problem in problems:
        print(problem)
    if check:
        return 1 if problems else 0
    if problems:
        raise HandoffError("fix the problems above before unpacking")
    if output is None or (output.exists() and any(output.iterdir())):
        raise HandoffError("unpack needs a new or empty --output-dir")
    output.mkdir(parents=True, exist_ok=True)
    for filename, fields, rows, _ in jobs:
        path = output / filename
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
        receipt["csv"][filename] = sha256(path)
    (output / "unpack_receipt.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(receipt, ensure_ascii=False, indent=2))
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    commands = parser.add_subparsers(dest="command", required=True)
    build_parser = commands.add_parser("build")
    build_parser.add_argument("--source-root", type=Path, required=True)
    build_parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    build_parser.add_argument("--audit-dir", type=Path, default=AUDIT_DIR)
    build_parser.add_argument("--part2-only", action="store_true")
    unpack_parser = commands.add_parser("unpack")
    unpack_parser.add_argument("--part1", type=Path)
    unpack_parser.add_argument("--part2", type=Path)
    unpack_parser.add_argument("--output-dir", type=Path)
    unpack_parser.add_argument("--check", action="store_true")
    unpack_parser.add_argument("--audit-dir", type=Path, default=AUDIT_DIR)
    args = parser.parse_args(argv)
    try:
        if args.command == "build":
            build(args.source_root, args.output_dir, args.audit_dir, args.part2_only)
            print(f"handoff folder: {args.output_dir}")
            return 0
        return unpack(args.part1, args.part2, args.output_dir, args.check,
                      args.audit_dir)
    except HandoffError as error:
        print(f"handoff error: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
